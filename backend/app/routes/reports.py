from datetime import date, timedelta
from decimal import Decimal, ROUND_HALF_UP
from io import StringIO, BytesIO
import csv

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session

from openpyxl import Workbook  # type: ignore[reportMissingModuleSource]
from openpyxl.styles import Font, Alignment  # type: ignore[reportMissingModuleSource]
from openpyxl.utils import get_column_letter  # type: ignore[reportMissingModuleSource]

from app.core.database import get_db
from app.core.security import get_current_user
from app.models import (
    Loan,
    Customer,
    Installment,
    Payment,
    User,
)


router = APIRouter(
    prefix="/reports",
    tags=["Reports"],
)


# =========================================================
# MONEY HELPER
# =========================================================

def money(value) -> Decimal:
    """
    Convert a numeric value into a safe money Decimal
    rounded to 2 decimal places.
    """

    return Decimal(
        str(value or 0)
    ).quantize(
        Decimal("0.01"),
        rounding=ROUND_HALF_UP,
    )


# =========================================================
# LOAN PORTFOLIO REPORT
# =========================================================

@router.get("/loan-portfolio")
def get_loan_portfolio_report(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    loans = (
        db.query(Loan)
        .order_by(Loan.id.desc())
        .all()
    )

    total_loans = len(loans)

    active_loans = sum(
        1
        for loan in loans
        if str(loan.status or "").lower() == "active"
    )

    completed_loans = sum(
        1
        for loan in loans
        if str(loan.status or "").lower() == "completed"
    )

    cancelled_loans = sum(
        1
        for loan in loans
        if str(loan.status or "").lower() == "cancelled"
    )

    total_principal = Decimal("0.00")
    total_interest = Decimal("0.00")
    total_payable = Decimal("0.00")
    total_collected = Decimal("0.00")
    total_outstanding = Decimal("0.00")

    loan_details = []

    for loan in loans:

        principal = money(
            loan.principal_amount
        )

        installment_totals = (
            db.query(
                func.coalesce(
                    func.sum(
                        Installment.interest_amount
                    ),
                    0,
                ),
                func.coalesce(
                    func.sum(
                        Installment.total_amount
                    ),
                    0,
                ),
            )
            .filter(
                Installment.loan_id == loan.id
            )
            .first()
        )

        interest = money(
            installment_totals[0]
        )

        installment_payable = money(
            installment_totals[1]
        )

        payments_collected = money(
            db.query(
                func.coalesce(
                    func.sum(Payment.amount),
                    0,
                )
            )
            .filter(
                Payment.loan_id == loan.id
            )
            .scalar()
        )

        if installment_payable > Decimal("0.00"):
            payable = installment_payable
        else:
            payable = money(
                principal + interest
            )

        outstanding = max(
            payable - payments_collected,
            Decimal("0.00"),
        )

        if payable > Decimal("0.00"):
            repayment_percentage = (
                payments_collected
                / payable
            ) * Decimal("100")

            repayment_percentage = min(
                max(
                    repayment_percentage,
                    Decimal("0.00"),
                ),
                Decimal("100.00"),
            )
        else:
            repayment_percentage = Decimal("0.00")

        customer = (
            db.query(Customer)
            .filter(
                Customer.id == loan.customer_id
            )
            .first()
        )

        customer_name = (
            customer.full_name
            if customer
            else "Unknown Customer"
        )

        total_installments = (
            db.query(
                func.count(Installment.id)
            )
            .filter(
                Installment.loan_id == loan.id
            )
            .scalar()
            or 0
        )

        paid_installments = (
            db.query(
                func.count(Installment.id)
            )
            .filter(
                Installment.loan_id == loan.id,
                Installment.paid_amount
                >= Installment.total_amount,
            )
            .scalar()
            or 0
        )

        unpaid_installments = max(
            total_installments
            - paid_installments,
            0,
        )

        total_principal += principal
        total_interest += interest
        total_payable += payable
        total_collected += payments_collected
        total_outstanding += outstanding

        loan_details.append(
            {
                "loan_id": loan.id,
                "customer_id": loan.customer_id,
                "customer_name": customer_name,
                "principal_amount": principal,
                "interest_rate": money(
                    loan.interest_rate
                ),
                "interest_type": loan.interest_type,
                "tenure_months": loan.tenure_months,
                "start_date": loan.start_date,
                "maturity_date": loan.maturity_date,
                "status": loan.status,
                "total_installments": total_installments,
                "paid_installments": paid_installments,
                "unpaid_installments": unpaid_installments,
                "total_interest": interest,
                "total_payable": payable,
                "total_collected": payments_collected,
                "outstanding": outstanding,
                "repayment_percentage": money(
                    repayment_percentage
                ),
            }
        )

    average_loan = (
        total_principal / total_loans
        if total_loans > 0
        else Decimal("0.00")
    )

    if total_payable > Decimal("0.00"):
        repayment_percentage = (
            total_collected
            / total_payable
        ) * Decimal("100")
    else:
        repayment_percentage = Decimal("0.00")

    repayment_percentage = min(
        max(
            repayment_percentage,
            Decimal("0.00"),
        ),
        Decimal("100.00"),
    )

    return {
        "summary": {
            "total_loans": total_loans,
            "active_loans": active_loans,
            "completed_loans": completed_loans,
            "cancelled_loans": cancelled_loans,
            "total_principal": money(
                total_principal
            ),
            "total_interest": money(
                total_interest
            ),
            "total_payable": money(
                total_payable
            ),
            "total_collected": money(
                total_collected
            ),
            "total_outstanding": money(
                total_outstanding
            ),
            "average_loan": money(
                average_loan
            ),
            "repayment_percentage": money(
                repayment_percentage
            ),
        },
        "loans": loan_details,
    }


# =========================================================
# PAYMENT COLLECTION REPORT
# =========================================================

@router.get("/payment-collection")
def get_payment_collection_report(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    payments = (
        db.query(Payment)
        .order_by(
            Payment.payment_date.desc(),
            Payment.id.desc(),
        )
        .all()
    )

    total_payments = len(payments)

    total_collected = Decimal("0.00")
    cash_collected = Decimal("0.00")
    upi_collected = Decimal("0.00")
    bank_collected = Decimal("0.00")
    other_collected = Decimal("0.00")

    cash_count = 0
    upi_count = 0
    bank_count = 0
    other_count = 0

    payment_details = []

    for payment in payments:

        amount = money(
            payment.amount
        )

        total_collected += amount

        payment_method = str(
            payment.payment_method or "other"
        ).strip().lower()

        if payment_method == "cash":
            cash_collected += amount
            cash_count += 1

        elif payment_method == "upi":
            upi_collected += amount
            upi_count += 1

        elif payment_method in {
            "bank",
            "bank transfer",
            "bank_transfer",
            "neft",
            "rtgs",
            "imps",
        }:
            bank_collected += amount
            bank_count += 1

        else:
            other_collected += amount
            other_count += 1

        loan = (
            db.query(Loan)
            .filter(
                Loan.id == payment.loan_id
            )
            .first()
        )

        customer = None

        if loan:
            customer = (
                db.query(Customer)
                .filter(
                    Customer.id == loan.customer_id
                )
                .first()
            )

        payment_details.append(
            {
                "payment_id": payment.id,
                "loan_id": payment.loan_id,
                "customer_id": (
                    customer.id
                    if customer
                    else None
                ),
                "customer_name": (
                    customer.full_name
                    if customer
                    else "Unknown Customer"
                ),
                "installment_id": payment.installment_id,
                "amount": amount,
                "payment_date": payment.payment_date,
                "payment_method": (
                    payment.payment_method
                    or "other"
                ),
                "reference_number": (
                    payment.reference_number
                ),
                "notes": payment.notes,
            }
        )

    total_collected = money(
        total_collected
    )

    cash_collected = money(
        cash_collected
    )

    upi_collected = money(
        upi_collected
    )

    bank_collected = money(
        bank_collected
    )

    other_collected = money(
        other_collected
    )

    average_payment = (
        total_collected / total_payments
        if total_payments > 0
        else Decimal("0.00")
    )

    return {
        "summary": {
            "total_payments": total_payments,
            "total_collected": total_collected,
            "average_payment": money(
                average_payment
            ),
            "cash_collected": cash_collected,
            "upi_collected": upi_collected,
            "bank_collected": bank_collected,
            "other_collected": other_collected,
            "cash_count": cash_count,
            "upi_count": upi_count,
            "bank_count": bank_count,
            "other_count": other_count,
        },
        "payments": payment_details,
    }


# =========================================================
# OUTSTANDING & OVERDUE REPORT
# =========================================================

@router.get("/outstanding-overdue")
def get_outstanding_overdue_report(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    today = date.today()
    next_7_days = today + timedelta(days=7)

    zero = Decimal("0.00")

    installments = (
        db.query(Installment)
        .join(
            Loan,
            Installment.loan_id == Loan.id
        )
        .filter(
            Loan.status == "active",
            Installment.paid_amount
            < Installment.total_amount,
        )
        .order_by(
            Installment.due_date.asc(),
            Installment.id.asc(),
        )
        .all()
    )

    total_outstanding = zero
    total_overdue = zero
    total_due_soon = zero

    overdue_count = 0
    due_soon_count = 0

    installment_details = []

    for installment in installments:

        total_amount = money(
            installment.total_amount
        )

        paid_amount = money(
            installment.paid_amount
        )

        outstanding = max(
            total_amount - paid_amount,
            zero,
        )

        loan = (
            db.query(Loan)
            .filter(
                Loan.id == installment.loan_id
            )
            .first()
        )

        if not loan:
            continue

        customer = (
            db.query(Customer)
            .filter(
                Customer.id == loan.customer_id
            )
            .first()
        )

        customer_name = (
            customer.full_name
            if customer
            else "Unknown Customer"
        )

        due_date = installment.due_date

        if due_date < today:

            status = "overdue"

            overdue_count += 1

            total_overdue += outstanding

        elif today <= due_date <= next_7_days:

            status = "due_soon"

            due_soon_count += 1

            total_due_soon += outstanding

        else:

            status = "upcoming"

        total_outstanding += outstanding

        installment_details.append(
            {
                "installment_id": installment.id,
                "loan_id": installment.loan_id,
                "customer_id": loan.customer_id,
                "customer_name": customer_name,
                "installment_number": (
                    installment.installment_number
                ),
                "due_date": due_date,
                "principal_amount": money(
                    installment.principal_amount
                ),
                "interest_amount": money(
                    installment.interest_amount
                ),
                "total_amount": total_amount,
                "paid_amount": paid_amount,
                "outstanding": outstanding,
                "status": status,
                "days_overdue": (
                    (today - due_date).days
                    if due_date < today
                    else 0
                ),
                "days_until_due": (
                    (due_date - today).days
                    if due_date >= today
                    else 0
                ),
            }
        )

    customer_summary_map = {}

    for item in installment_details:

        customer_id = item["customer_id"]

        if customer_id not in customer_summary_map:

            customer_summary_map[
                customer_id
            ] = {
                "customer_id": customer_id,
                "customer_name": item[
                    "customer_name"
                ],
                "outstanding": zero,
                "overdue_amount": zero,
                "due_soon_amount": zero,
                "installments": 0,
                "overdue_installments": 0,
            }

        customer_data = customer_summary_map[
            customer_id
        ]

        customer_data["outstanding"] += item[
            "outstanding"
        ]

        customer_data["installments"] += 1

        if item["status"] == "overdue":

            customer_data[
                "overdue_amount"
            ] += item["outstanding"]

            customer_data[
                "overdue_installments"
            ] += 1

        elif item["status"] == "due_soon":

            customer_data[
                "due_soon_amount"
            ] += item["outstanding"]

    customer_summary = []

    for customer_data in customer_summary_map.values():

        customer_summary.append(
            {
                "customer_id": customer_data[
                    "customer_id"
                ],
                "customer_name": customer_data[
                    "customer_name"
                ],
                "outstanding": money(
                    customer_data["outstanding"]
                ),
                "overdue_amount": money(
                    customer_data["overdue_amount"]
                ),
                "due_soon_amount": money(
                    customer_data["due_soon_amount"]
                ),
                "installments": customer_data[
                    "installments"
                ],
                "overdue_installments": customer_data[
                    "overdue_installments"
                ],
            }
        )

    customer_summary.sort(
        key=lambda item: item["outstanding"],
        reverse=True,
    )

    return {
        "summary": {
            "total_outstanding": money(
                total_outstanding
            ),
            "total_overdue": money(
                total_overdue
            ),
            "total_due_soon": money(
                total_due_soon
            ),
            "overdue_count": overdue_count,
            "due_soon_count": due_soon_count,
            "total_open_installments": len(
                installment_details
            ),
        },
        "installments": installment_details,
        "customers": customer_summary,
    }


# =========================================================
# CUSTOMER FINANCIAL REPORT
# =========================================================

@router.get("/customer-financial")
def get_customer_financial_report(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    customers = (
        db.query(Customer)
        .order_by(Customer.full_name.asc())
        .all()
    )

    customer_reports = []

    overall_principal = Decimal("0.00")
    overall_payable = Decimal("0.00")
    overall_paid = Decimal("0.00")
    overall_outstanding = Decimal("0.00")

    customers_with_loans = 0

    for customer in customers:

        loans = (
            db.query(Loan)
            .filter(
                Loan.customer_id == customer.id
            )
            .order_by(Loan.id.asc())
            .all()
        )

        if loans:
            customers_with_loans += 1

        total_principal = Decimal("0.00")
        total_interest = Decimal("0.00")
        total_payable = Decimal("0.00")
        total_paid = Decimal("0.00")

        payment_count = 0

        active_loans = 0
        completed_loans = 0
        cancelled_loans = 0

        loan_history = []

        for loan in loans:

            if loan.status == "active":
                active_loans += 1

            elif loan.status == "completed":
                completed_loans += 1

            elif loan.status == "cancelled":
                cancelled_loans += 1

            installments = (
                db.query(Installment)
                .filter(
                    Installment.loan_id == loan.id
                )
                .order_by(
                    Installment.installment_number.asc()
                )
                .all()
            )

            payments = (
                db.query(Payment)
                .filter(
                    Payment.loan_id == loan.id
                )
                .all()
            )

            loan_principal = money(
                loan.principal_amount
            )

            loan_interest = sum(
                (
                    money(
                        installment.interest_amount
                    )
                    for installment in installments
                ),
                Decimal("0.00"),
            )

            if installments:

                loan_payable = sum(
                    (
                        money(
                            installment.total_amount
                        )
                        for installment in installments
                    ),
                    Decimal("0.00"),
                )

            else:

                loan_payable = (
                    loan_principal
                    + loan_interest
                )

            loan_paid = sum(
                (
                    money(payment.amount)
                    for payment in payments
                ),
                Decimal("0.00"),
            )

            loan_outstanding = max(
                loan_payable - loan_paid,
                Decimal("0.00"),
            )

            if loan_payable > Decimal("0.00"):

                loan_repayment_percentage = (
                    loan_paid
                    / loan_payable
                    * Decimal("100")
                )

                loan_repayment_percentage = min(
                    max(
                        loan_repayment_percentage,
                        Decimal("0.00"),
                    ),
                    Decimal("100.00"),
                )

                loan_repayment_percentage = (
                    loan_repayment_percentage.quantize(
                        Decimal("0.01"),
                        rounding=ROUND_HALF_UP,
                    )
                )

            else:

                loan_repayment_percentage = (
                    Decimal("0.00")
                )

            paid_installments = sum(
                1
                for installment in installments
                if money(
                    installment.paid_amount
                )
                >= money(
                    installment.total_amount
                )
            )

            unpaid_installments = (
                len(installments)
                - paid_installments
            )

            total_principal += loan_principal
            total_interest += loan_interest
            total_payable += loan_payable
            total_paid += loan_paid
            payment_count += len(payments)

            loan_history.append(
                {
                    "loan_id": loan.id,
                    "principal_amount": loan_principal,
                    "interest_rate": money(
                        loan.interest_rate
                    ),
                    "interest_type": loan.interest_type,
                    "tenure_months": loan.tenure_months,
                    "start_date": loan.start_date,
                    "maturity_date": loan.maturity_date,
                    "status": loan.status,
                    "total_installments": len(
                        installments
                    ),
                    "paid_installments": paid_installments,
                    "unpaid_installments": (
                        unpaid_installments
                    ),
                    "total_interest": money(
                        loan_interest
                    ),
                    "total_payable": money(
                        loan_payable
                    ),
                    "total_paid": money(
                        loan_paid
                    ),
                    "outstanding": money(
                        loan_outstanding
                    ),
                    "repayment_percentage": (
                        loan_repayment_percentage
                    ),
                }
            )

        outstanding = max(
            total_payable - total_paid,
            Decimal("0.00"),
        )

        if total_payable > Decimal("0.00"):

            repayment_percentage = (
                total_paid
                / total_payable
                * Decimal("100")
            )

            repayment_percentage = min(
                max(
                    repayment_percentage,
                    Decimal("0.00"),
                ),
                Decimal("100.00"),
            )

            repayment_percentage = (
                repayment_percentage.quantize(
                    Decimal("0.01"),
                    rounding=ROUND_HALF_UP,
                )
            )

        else:

            repayment_percentage = (
                Decimal("0.00")
            )

        total_principal = money(
            total_principal
        )

        total_interest = money(
            total_interest
        )

        total_payable = money(
            total_payable
        )

        total_paid = money(
            total_paid
        )

        outstanding = money(
            outstanding
        )

        overall_principal += total_principal
        overall_payable += total_payable
        overall_paid += total_paid
        overall_outstanding += outstanding

        customer_reports.append(
            {
                "customer_id": customer.id,
                "customer_name": customer.full_name,
                "total_loans": len(loans),
                "active_loans": active_loans,
                "completed_loans": completed_loans,
                "cancelled_loans": cancelled_loans,
                "total_principal": total_principal,
                "total_interest": total_interest,
                "total_payable": total_payable,
                "total_paid": total_paid,
                "outstanding": outstanding,
                "payment_count": payment_count,
                "repayment_percentage": (
                    repayment_percentage
                ),
                "loans": loan_history,
            }
        )

    overall_principal = money(
        overall_principal
    )

    overall_payable = money(
        overall_payable
    )

    overall_paid = money(
        overall_paid
    )

    overall_outstanding = money(
        overall_outstanding
    )

    if overall_payable > Decimal("0.00"):

        overall_repayment_percentage = (
            overall_paid
            / overall_payable
            * Decimal("100")
        )

        overall_repayment_percentage = min(
            max(
                overall_repayment_percentage,
                Decimal("0.00"),
            ),
            Decimal("100.00"),
        )

        overall_repayment_percentage = (
            overall_repayment_percentage.quantize(
                Decimal("0.01"),
                rounding=ROUND_HALF_UP,
            )
        )

    else:

        overall_repayment_percentage = (
            Decimal("0.00")
        )

    return {
        "summary": {
            "total_customers": len(customers),
            "customers_with_loans": (
                customers_with_loans
            ),
            "total_principal": overall_principal,
            "total_payable": overall_payable,
            "total_collected": overall_paid,
            "total_outstanding": (
                overall_outstanding
            ),
            "repayment_percentage": (
                overall_repayment_percentage
            ),
        },
        "customers": customer_reports,
    }


# =========================================================
# DATE RANGE REPORT
# =========================================================

@router.get("/date-range")
def get_date_range_report(
    from_date: date = Query(...),
    to_date: date = Query(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if from_date > to_date:
        return {
            "error": (
                "from_date cannot be later than to_date"
            )
        }

    payments = (
        db.query(Payment)
        .filter(
            Payment.payment_date >= from_date,
            Payment.payment_date <= to_date,
        )
        .order_by(
            Payment.payment_date.desc()
        )
        .all()
    )

    total_collected = Decimal("0.00")

    payment_records = []

    for payment in payments:

        amount = money(
            payment.amount
        )

        total_collected += amount

        customer_name = "Unknown"

        if payment.loan and payment.loan.customer:

            customer_name = (
                payment.loan.customer.full_name
            )

        payment_records.append(
            {
                "payment_id": payment.id,
                "loan_id": payment.loan_id,
                "customer_id": (
                    payment.loan.customer_id
                    if payment.loan
                    else None
                ),
                "customer_name": customer_name,
                "amount": amount,
                "payment_date": payment.payment_date,
                "payment_method": (
                    payment.payment_method
                ),
                "reference_number": (
                    payment.reference_number
                ),
            }
        )

    loans = (
        db.query(Loan)
        .filter(
            Loan.start_date >= from_date,
            Loan.start_date <= to_date,
        )
        .order_by(
            Loan.start_date.desc()
        )
        .all()
    )

    total_principal = Decimal("0.00")

    loan_records = []

    for loan in loans:

        principal = money(
            loan.principal_amount
        )

        total_principal += principal

        customer_name = "Unknown"

        if loan.customer:

            customer_name = (
                loan.customer.full_name
            )

        loan_records.append(
            {
                "loan_id": loan.id,
                "customer_id": loan.customer_id,
                "customer_name": customer_name,
                "principal_amount": principal,
                "interest_rate": loan.interest_rate,
                "interest_type": loan.interest_type,
                "tenure_months": loan.tenure_months,
                "start_date": loan.start_date,
                "maturity_date": loan.maturity_date,
                "status": loan.status,
            }
        )

    cash_collected = Decimal("0.00")
    upi_collected = Decimal("0.00")
    bank_collected = Decimal("0.00")
    other_collected = Decimal("0.00")

    cash_count = 0
    upi_count = 0
    bank_count = 0
    other_count = 0

    for payment in payments:

        amount = money(
            payment.amount
        )

        method = str(
            payment.payment_method or ""
        ).strip().lower()

        if method == "cash":

            cash_collected += amount
            cash_count += 1

        elif method == "upi":

            upi_collected += amount
            upi_count += 1

        elif method in {
            "bank",
            "bank transfer",
            "bank_transfer",
            "neft",
            "rtgs",
            "imps",
        }:

            bank_collected += amount
            bank_count += 1

        else:

            other_collected += amount
            other_count += 1

    average_payment = (
        total_collected / len(payments)
        if payments
        else Decimal("0.00")
    )

    return {
        "date_range": {
            "from_date": from_date,
            "to_date": to_date,
        },
        "summary": {
            "total_loans": len(loans),
            "total_principal": money(
                total_principal
            ),
            "total_payments": len(payments),
            "total_collected": money(
                total_collected
            ),
            "average_payment": money(
                average_payment
            ),
            "cash_collected": money(
                cash_collected
            ),
            "upi_collected": money(
                upi_collected
            ),
            "bank_collected": money(
                bank_collected
            ),
            "other_collected": money(
                other_collected
            ),
            "cash_count": cash_count,
            "upi_count": upi_count,
            "bank_count": bank_count,
            "other_count": other_count,
        },
        "loans": loan_records,
        "payments": payment_records,
    }


# =========================================================
# EXPORT FOUNDATION - CSV
# =========================================================

def create_csv_response(
    filename: str,
    headers: list[str],
    rows: list[list],
):
    output = StringIO()

    writer = csv.writer(output)

    writer.writerow(headers)

    for row in rows:
        writer.writerow(row)

    output.seek(0)

    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )


# =========================================================
# EXPORT FOUNDATION - EXCEL
# =========================================================

def create_excel_response(
    filename: str,
    sheet_name: str,
    headers: list[str],
    rows: list[list],
):
    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = sheet_name

    # -----------------------------------------------------
    # HEADER ROW
    # -----------------------------------------------------

    for column_number, header in enumerate(
        headers,
        start=1,
    ):

        cell = worksheet.cell(
            row=1,
            column=column_number,
            value=header,
        )

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # -----------------------------------------------------
    # DATA ROWS
    # -----------------------------------------------------

    for row_number, row in enumerate(
        rows,
        start=2,
    ):

        for column_number, value in enumerate(
            row,
            start=1,
        ):

            cell = worksheet.cell(
                row=row_number,
                column=column_number,
                value=value,
            )

            cell.alignment = Alignment(
                vertical="center",
            )

    # -----------------------------------------------------
    # FREEZE HEADER
    # -----------------------------------------------------

    worksheet.freeze_panes = "A2"

    # -----------------------------------------------------
    # AUTO-SIZE COLUMNS
    # -----------------------------------------------------

    for column_number, header in enumerate(
        headers,
        start=1,
    ):

        column_letter = get_column_letter(
            column_number
        )

        max_length = len(
            str(header)
        )

        for row in rows:

            if column_number <= len(row):

                value = row[
                    column_number - 1
                ]

                if value is not None:

                    max_length = max(
                        max_length,
                        len(str(value)),
                    )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 3,
            35,
        )

    # -----------------------------------------------------
    # KEEP STRING VALUES AS TEXT
    # -----------------------------------------------------

    for row in worksheet.iter_rows(
        min_row=2
    ):

        for cell in row:

            if isinstance(
                cell.value,
                str,
            ):

                cell.number_format = "@"

    # -----------------------------------------------------
    # EXCEL DATE FORMATTING
    # -----------------------------------------------------

    for row in worksheet.iter_rows(
        min_row=2
    ):

        for cell in row:

            if isinstance(
                cell.value,
                date,
            ):

                cell.number_format = "dd-mm-yyyy"

    # -----------------------------------------------------
    # EXCEL FILE IN MEMORY
    # -----------------------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )


# =========================================================
# EXPORT LOAN PORTFOLIO TO CSV
# =========================================================

@router.get("/export/loan-portfolio")
def export_loan_portfolio(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    loans = (
        db.query(Loan)
        .order_by(Loan.id.desc())
        .all()
    )

    payments = db.query(
        Payment
    ).all()

    payments_by_loan = {}

    for payment in payments:

        loan_id = payment.loan_id

        if loan_id not in payments_by_loan:

            payments_by_loan[
                loan_id
            ] = Decimal("0.00")

        payments_by_loan[
            loan_id
        ] += money(
            payment.amount
        )

    rows = []

    for loan in loans:

        installments = (
            loan.installments or []
        )

        principal = money(
            loan.principal_amount
        )

        total_interest = Decimal(
            "0.00"
        )

        total_payable = Decimal(
            "0.00"
        )

        paid_installments = 0

        for installment in installments:

            interest_amount = money(
                installment.interest_amount
            )

            total_amount = money(
                installment.total_amount
            )

            paid_amount = money(
                installment.paid_amount
            )

            total_interest += (
                interest_amount
            )

            total_payable += (
                total_amount
            )

            if paid_amount >= total_amount:

                paid_installments += 1

        if not installments:

            total_payable = principal

        total_collected = (
            payments_by_loan.get(
                loan.id,
                Decimal("0.00"),
            )
        )

        outstanding = (
            total_payable
            - total_collected
        )

        if outstanding < 0:

            outstanding = Decimal(
                "0.00"
            )

        total_installments = len(
            installments
        )

        unpaid_installments = (
            total_installments
            - paid_installments
        )

        repayment_percentage = (
            Decimal("0.00")
        )

        if total_payable > 0:

            repayment_percentage = (
                total_collected
                / total_payable
            ) * Decimal("100")

        customer_name = "Unknown"

        if loan.customer:

            customer_name = (
                loan.customer.full_name
            )

        rows.append(
            [
                loan.id,
                loan.customer_id,
                customer_name,
                money(principal),
                money(loan.interest_rate),
                loan.interest_type,
                loan.tenure_months,
                loan.start_date,
                loan.maturity_date,
                loan.status,
                total_installments,
                paid_installments,
                unpaid_installments,
                money(total_interest),
                money(total_payable),
                money(total_collected),
                money(outstanding),
                money(
                    repayment_percentage
                ),
            ]
        )

    headers = [
        "Loan ID",
        "Customer ID",
        "Customer Name",
        "Principal Amount",
        "Interest Rate",
        "Interest Type",
        "Tenure Months",
        "Start Date",
        "Maturity Date",
        "Status",
        "Total Installments",
        "Paid Installments",
        "Unpaid Installments",
        "Total Interest",
        "Total Payable",
        "Total Collected",
        "Outstanding",
        "Repayment Percentage",
    ]

    return create_csv_response(
        filename="loan_portfolio_report.csv",
        headers=headers,
        rows=rows,
    )


# =========================================================
# EXPORT PAYMENT COLLECTION TO CSV
# =========================================================

@router.get("/export/payment-collection")
def export_payment_collection(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    payments = (
        db.query(Payment)
        .order_by(
            Payment.payment_date.desc(),
            Payment.id.desc(),
        )
        .all()
    )

    rows = []

    for payment in payments:

        customer_id = None
        customer_name = "Unknown"

        if (
            payment.loan
            and payment.loan.customer
        ):

            customer_id = (
                payment.loan.customer_id
            )

            customer_name = (
                payment.loan.customer.full_name
            )

        rows.append(
            [
                payment.id,
                payment.loan_id,
                customer_id,
                customer_name,
                payment.installment_id,
                money(payment.amount),
                payment.payment_date,
                payment.payment_method,
                payment.reference_number,
                payment.notes,
            ]
        )

    headers = [
        "Payment ID",
        "Loan ID",
        "Customer ID",
        "Customer Name",
        "Installment ID",
        "Amount",
        "Payment Date",
        "Payment Method",
        "Reference Number",
        "Notes",
    ]

    return create_csv_response(
        filename="payment_collection_report.csv",
        headers=headers,
        rows=rows,
    )


# =========================================================
# EXPORT OUTSTANDING & OVERDUE TO CSV
# =========================================================

@router.get("/export/outstanding-overdue")
def export_outstanding_overdue(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    today = date.today()

    unpaid_installments = (
        db.query(Installment)
        .filter(
            Installment.paid_amount
            < Installment.total_amount
        )
        .order_by(
            Installment.due_date.asc()
        )
        .all()
    )

    rows = []

    for installment in unpaid_installments:

        customer_id = None
        customer_name = "Unknown"

        loan_id = installment.loan_id

        if installment.loan:

            customer_id = (
                installment.loan.customer_id
            )

            if installment.loan.customer:

                customer_name = (
                    installment.loan.customer.full_name
                )

        paid_amount = money(
            installment.paid_amount
        )

        total_amount = money(
            installment.total_amount
        )

        outstanding = (
            total_amount
            - paid_amount
        )

        if outstanding < 0:

            outstanding = Decimal(
                "0.00"
            )

        due_date = (
            installment.due_date
        )

        status = "Upcoming"

        days_overdue = 0
        days_until_due = 0

        if due_date < today:

            status = "Overdue"

            days_overdue = (
                today - due_date
            ).days

        elif due_date == today:

            status = "Due Today"

        else:

            days_until_due = (
                due_date - today
            ).days

            if days_until_due <= 7:

                status = "Due Soon"

        rows.append(
            [
                installment.id,
                loan_id,
                customer_id,
                customer_name,
                installment.installment_number,
                due_date,
                money(
                    installment.principal_amount
                ),
                money(
                    installment.interest_amount
                ),
                total_amount,
                paid_amount,
                money(outstanding),
                status,
                days_overdue,
                days_until_due,
            ]
        )

    headers = [
        "Installment ID",
        "Loan ID",
        "Customer ID",
        "Customer Name",
        "Installment Number",
        "Due Date",
        "Principal Amount",
        "Interest Amount",
        "Total Amount",
        "Paid Amount",
        "Outstanding",
        "Status",
        "Days Overdue",
        "Days Until Due",
    ]

    return create_csv_response(
        filename="outstanding_overdue_report.csv",
        headers=headers,
        rows=rows,
    )


# =========================================================
# EXPORT CUSTOMER FINANCIAL TO CSV
# =========================================================

@router.get("/export/customer-financial")
def export_customer_financial(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    customers = (
        db.query(Customer)
        .order_by(Customer.id.asc())
        .all()
    )

    rows = []

    for customer in customers:

        loans = (
            db.query(Loan)
            .filter(
                Loan.customer_id == customer.id
            )
            .all()
        )

        total_principal = Decimal(
            "0.00"
        )

        total_interest = Decimal(
            "0.00"
        )

        total_payable = Decimal(
            "0.00"
        )

        total_paid = Decimal(
            "0.00"
        )

        active_loans = 0
        completed_loans = 0
        cancelled_loans = 0

        payment_count = 0

        for loan in loans:

            total_principal += money(
                loan.principal_amount
            )

            if loan.status == "active":

                active_loans += 1

            elif loan.status == "completed":

                completed_loans += 1

            elif loan.status == "cancelled":

                cancelled_loans += 1

            installments = (
                loan.installments or []
            )

            for installment in installments:

                total_interest += money(
                    installment.interest_amount
                )

                total_payable += money(
                    installment.total_amount
                )

            loan_payments = (
                db.query(Payment)
                .filter(
                    Payment.loan_id == loan.id
                )
                .all()
            )

            for payment in loan_payments:

                total_paid += money(
                    payment.amount
                )

                payment_count += 1

            if not installments:

                total_payable += money(
                    loan.principal_amount
                )

        outstanding = (
            total_payable
            - total_paid
        )

        if outstanding < 0:

            outstanding = Decimal(
                "0.00"
            )

        repayment_percentage = (
            Decimal("0.00")
        )

        if total_payable > 0:

            repayment_percentage = (
                total_paid
                / total_payable
            ) * Decimal("100")

        rows.append(
            [
                customer.id,
                customer.full_name,
                len(loans),
                active_loans,
                completed_loans,
                cancelled_loans,
                money(total_principal),
                money(total_interest),
                money(total_payable),
                money(total_paid),
                money(outstanding),
                payment_count,
                money(
                    repayment_percentage
                ),
            ]
        )

    headers = [
        "Customer ID",
        "Customer Name",
        "Total Loans",
        "Active Loans",
        "Completed Loans",
        "Cancelled Loans",
        "Total Principal",
        "Total Interest",
        "Total Payable",
        "Total Paid",
        "Outstanding",
        "Payment Count",
        "Repayment Percentage",
    ]

    return create_csv_response(
        filename="customer_financial_report.csv",
        headers=headers,
        rows=rows,
    )


# =========================================================
# TEST EXCEL EXPORT
# =========================================================

@router.get("/export/test-excel")
def export_test_excel(
    current_user: User = Depends(get_current_user),
):
    headers = [
        "Test ID",
        "Name",
        "Amount",
        "Date",
        "Reference",
    ]

    rows = [
        [
            1,
            "Excel Export Test",
            12345.67,
            date.today(),
            "TEST-123456789012345",
        ]
    ]

    return create_excel_response(
        filename="finance_loan_excel_test.xlsx",
        sheet_name="Export Test",
        headers=headers,
        rows=rows,
    )